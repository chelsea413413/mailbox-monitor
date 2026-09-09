"""Excel detail attachment generator (openpyxl)."""
from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from .models import CrawlResult, LetterRecord, RecordStatus


HEADERS = [
    "省份", "状态", "标题", "发布日期", "回复日期",
    "咨询内容", "回复内容", "原文URL", "网站原始ID",
]

COLUMN_WIDTHS = [8, 8, 40, 14, 14, 60, 60, 45, 20]

STATUS_CN = {
    RecordStatus.NEW: "新增",
    RecordStatus.UPDATED: "更新",
    RecordStatus.UNCHANGED: "无变化",
}

HEADER_FILL = PatternFill(start_color="2B579A", end_color="2B579A", fill_type="solid")
HEADER_FONT = Font(color="FFFFFF", bold=True, size=11)
NEW_FILL = PatternFill(start_color="E2EFDA", end_color="E2EFDA", fill_type="solid")
UPDATED_FILL = PatternFill(start_color="FFF2CC", end_color="FFF2CC", fill_type="solid")


def generate_excel(
    results: list[CrawlResult],
    output_path: str,
) -> str:
    """Generate an Excel file with all changed records across provinces.

    Only NEW and UPDATED records are included. Each province gets its own
    sheet. If a province has no changes, it is omitted.
    """
    wb = Workbook()
    wb.remove(wb.active)

    any_data = False

    for result in results:
        changed = result.changed_records
        if not changed:
            continue
        any_data = True

        ws = wb.create_sheet(title=result.province[:31])
        _write_sheet(ws, changed)

    if not any_data:
        ws = wb.create_sheet(title="今日无变化")
        ws["A1"] = "今日无新增或更新的信件"
        ws["A1"].font = Font(size=14, bold=True)
        ws.merge_cells("A1:C1")

    Path(output_path).parent.mkdir(parents=True, exist_ok=True)
    wb.save(output_path)
    return output_path


def _write_sheet(ws, records: list[LetterRecord]) -> None:
    """Write records into a worksheet with styling."""
    for col_idx, header in enumerate(HEADERS, 1):
        cell = ws.cell(row=1, column=col_idx, value=header)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(horizontal="center", vertical="center")
        ws.column_dimensions[get_column_letter(col_idx)].width = COLUMN_WIDTHS[col_idx - 1]

    for row_idx, record in enumerate(records, 2):
        values = [
            record.province,
            STATUS_CN.get(record.status, str(record.status)),
            record.title,
            record.publish_date or "",
            record.reply_date or "",
            record.content,
            record.reply_content,
            record.url,
            record.original_id,
        ]
        for col_idx, val in enumerate(values, 1):
            cell = ws.cell(row=row_idx, column=col_idx, value=val)
            cell.alignment = Alignment(vertical="top", wrap_text=True)

        if record.status == RecordStatus.NEW:
            fill = NEW_FILL
        elif record.status == RecordStatus.UPDATED:
            fill = UPDATED_FILL
        else:
            fill = None
        if fill:
            for col_idx in range(1, len(HEADERS) + 1):
                ws.cell(row=row_idx, column=col_idx).fill = fill

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(HEADERS))}{len(records) + 1}"
